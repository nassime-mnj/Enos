"""Tests de l'éditeur PDF (nasko_pdf). Lancer : python -m unittest tests.test_pdf -v"""
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pymupdf  # noqa: E402

import nasko_pdf as P  # noqa: E402


def make_pdf(path, texts, size=(595, 842)):
    d = pymupdf.open()
    for t in texts:
        p = d.new_page(width=size[0], height=size[1])
        p.insert_text((72, 100), t, fontsize=18)
    d.save(path)
    d.close()
    return path


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        P.ROOT = os.path.join(self.tmp, "Editeur PDF")
        P.WORK = os.path.join(P.ROOT, ".work")
        P.OUT = os.path.join(P.ROOT, "Livraison")
        P.S.update(name="", path="", hist=[], ver=0)
        P.JOB.update(busy=False, result=None, error="", cancel=False)
        P.ensure_dirs()
        self.a = make_pdf(os.path.join(self.tmp, "a.pdf"), ["Alpha un", "Alpha deux", "Alpha trois"])
        self.b = make_pdf(os.path.join(self.tmp, "b.pdf"), ["Bravo un", "Bravo deux"])

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def texts(self):
        return [t.strip() for t in P.page_texts()]

    def open_a(self):
        P.add_files([(self.a, "a.pdf")])


class Pages(Base):
    def test_merge_pdfs_and_image(self):
        from PIL import Image
        img = os.path.join(self.tmp, "x.png")
        Image.new("RGBA", (400, 300), (255, 0, 0, 128)).save(img)
        P.add_files([(self.a, "a.pdf"), (self.b, "b.pdf"), (img, "x.png")])
        st = P.state()
        self.assertEqual(len(st["pages"]), 6)
        self.assertEqual(st["name"], "fusion.pdf")
        self.assertGreater(st["pages"][5]["w"], st["pages"][5]["h"])          # image paysage -> page paysage
        self.assertEqual(self.texts()[:2], ["Alpha un", "Alpha deux"])

    def test_append_then_undo(self):
        self.open_a()
        P.add_files([(self.b, "b.pdf")])
        self.assertEqual(len(P.state()["pages"]), 5)
        self.assertEqual(len(P.undo()["pages"]), 3)

    def test_arrange_reorder_rotate_delete_duplicate_blank(self):
        self.open_a()
        P.arrange([{"src": 2, "rot": 90}, {"src": 0}, {"src": -1}, {"src": 0}])
        st = P.state()
        self.assertEqual(len(st["pages"]), 4)
        self.assertEqual([p["rot"] for p in st["pages"]], [90, 0, 0, 0])
        self.assertEqual(self.texts(), ["Alpha trois", "Alpha un", "", "Alpha un"])
        with self.assertRaises(P.PdfError):
            P.arrange([])

    def test_parse_pages(self):
        self.assertEqual(P.parse_pages("1-2, 5, 8-", 9), [0, 1, 4, 7, 8])
        self.assertEqual(P.parse_pages("impaires", 5), [0, 2, 4])
        self.assertEqual(P.parse_pages("", 3), [0, 1, 2])
        for bad in ("0", "4", "a-b", "3-9"):
            with self.assertRaises(P.PdfError):
                P.parse_pages(bad, 3)

    def test_extract_and_split(self):
        P.add_files([(self.a, "a.pdf"), (self.b, "b.pdf")])
        out = P.extract("2-3")
        d = pymupdf.open(out[0])
        self.assertEqual([p.get_text().strip() for p in d], ["Alpha deux", "Alpha trois"])
        files = P.split("every", "2")
        self.assertTrue(files[0].endswith(".zip"))
        self.assertEqual([pymupdf.open(f).page_count for f in files[1:]], [2, 2, 1])
        self.assertEqual(len(P.split("each")[1:]), 5)
        self.assertEqual([pymupdf.open(f).page_count for f in P.split("ranges", "1-2; 3-")[1:]], [2, 3])
        with self.assertRaises(P.PdfError):
            P.split("bookmarks")
        self.assertEqual(len(P.state()["pages"]), 5)                      # le document de travail n'est pas modifié

    def test_render(self):
        self.open_a()
        data, mime = P.render(0, 0.5)
        self.assertEqual(mime, "image/jpeg")
        self.assertTrue(data[:2] == b"\xff\xd8")
        with self.assertRaises(P.PdfError):
            P.render(9)


class Edit(Base):
    def test_redaction_is_irreversible_and_scrubs_metadata(self):
        d = pymupdf.open()
        p = d.new_page()
        p.insert_text((72, 100), "IBAN FR7630006000011234567890189 secret", fontsize=12)
        p.insert_text((72, 200), "Texte public", fontsize=12)
        d.set_metadata({"author": "Jean Dupont"})
        path = os.path.join(self.tmp, "r.pdf")
        d.save(path)
        iban_rect = d[0].search_for("FR7630006000011234567890189")[0]
        P.add_files([(path, "r.pdf")])
        P.annotate([{"type": "redact_text", "preset": "iban"}])
        txt = P.page_texts()[0]
        self.assertNotIn("FR76", txt)
        self.assertIn("Texte public", txt)
        self.assertIn("secret", txt)
        d2 = pymupdf.open(P.S["path"])
        self.assertFalse(d2.metadata.get("author"))
        pix = d2[0].get_pixmap(dpi=72, clip=iban_rect)
        self.assertLess(sum(pix.samples) / len(pix.samples), 40)           # zone noire
        with self.assertRaises(P.PdfError):
            P.annotate([{"type": "redact_text", "query": "introuvable-xyz"}])

    def test_redact_rect_removes_text(self):
        self.open_a()
        P.annotate([{"type": "redact", "page": 0, "rect": [0.05, 0.05, 0.6, 0.2]}])
        self.assertEqual(self.texts()[0], "")
        self.assertEqual(self.texts()[1], "Alpha deux")

    def test_text_lines_and_edittext(self):
        self.open_a()
        lines = P.text_lines(0)
        self.assertEqual(lines[0]["text"], "Alpha un")
        ln = lines[0]
        P.annotate([{"type": "edittext", "page": 0, "rect": [ln["x0"], ln["y0"], ln["x1"], ln["y1"]], "text": "Alpha CORRIGÉ", "size": ln["size"],
                     "color": ln["color"], "family": ln["family"], "old_len": len(ln["text"])}])
        t = self.texts()[0]
        self.assertIn("CORRIGÉ", t)
        self.assertNotIn("Alpha un", t)

    def test_edittext_on_rotated_page_keeps_line_on_one_row(self):
        self.open_a()
        P.arrange([{"src": 0, "rot": 90}])
        ln = P.text_lines(0)[0]
        self.assertEqual(ln["angle"], 0)                                   # le texte garde son sens dans le repère non pivoté
        P.annotate([{"type": "edittext", "page": 0, "rect": [ln["x0"], ln["y0"], ln["x1"], ln["y1"]], "text": "Alpha un, version très prolongée", "size": ln["size"],
                     "color": ln["color"], "family": ln["family"], "old_len": len(ln["text"]), "angle": ln["angle"]}])
        self.assertEqual(self.texts()[0], "Alpha un, version très prolongée")

    def test_annotations_on_rotated_page_stay_upright(self):
        self.open_a()
        P.arrange([{"src": 0, "rot": 90}, {"src": 1}])
        P.annotate([
            {"type": "text", "page": 0, "rect": [0.2, 0.3, 0.8, 0.4], "text": "Bonjour tourné", "size": 14, "color": "#ff0000"},
            {"type": "stamp", "page": 0, "rect": [0.1, 0.6, 0.5, 0.7], "label": "APPROUVÉ", "color": "#008800"},
            {"type": "highlight", "page": 1, "rect": [0.1, 0.1, 0.5, 0.2]},
            {"type": "note", "page": 1, "x": 0.5, "y": 0.5, "text": "À vérifier"},
            {"type": "ink", "page": 1, "points": [[0.1, 0.1], [0.3, 0.3], [0.5, 0.2]]},
        ])
        d = pymupdf.open(P.S["path"])
        page = d[0]
        hits = page.search_for("Bonjour")
        self.assertTrue(hits)
        r = hits[0] * page.rotation_matrix                                  # les rectangles de recherche sont dans le repère non pivoté
        self.assertGreater(r.width, r.height)                              # texte horizontal à l'affichage
        cx, cy = (r.x0 + r.x1) / 2 / page.rect.width, (r.y0 + r.y1) / 2 / page.rect.height
        self.assertTrue(0.2 < cx < 0.8 and 0.28 < cy < 0.42, (cx, cy))     # au bon endroit de la page affichée
        self.assertTrue(page.search_for("APPROUV"))
        self.assertEqual(len(list(d[1].annots())), 3)

    def test_signature_image(self):
        import base64
        import io
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGBA", (200, 80), (0, 0, 255, 255)).save(buf, "PNG")
        url = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
        self.open_a()
        P.annotate([{"type": "image", "page": 0, "rect": [0.5, 0.8, 0.9, 0.9], "data": url}])
        self.assertEqual(len(pymupdf.open(P.S["path"])[0].get_images()), 1)
        with self.assertRaises(P.PdfError):
            P.annotate([{"type": "image", "page": 0, "rect": [0, 0, 1, 1], "data": "data:text/html;base64,AAAA"}])

    def test_watermark_and_numbers(self):
        self.open_a()
        P.watermark({"text": "CONFIDENTIEL", "angle": 45, "opacity": 0.2, "pages": "1-2"})
        d = pymupdf.open(P.S["path"])
        self.assertTrue(d[0].search_for("CONFIDENTIEL"))
        self.assertFalse(d[2].search_for("CONFIDENTIEL"))
        dirv = [ln["dir"] for b in d[0].get_text("dict")["blocks"] if b.get("type") == 0 for ln in b["lines"] if "CONFIDENTIEL" in "".join(s["text"] for s in ln["spans"])][0]
        self.assertLess(dirv[1], -0.5)                                     # monte vers la droite
        P.page_numbers({"format": "Page {n}/{N}", "position": "bottom-center", "skip_first": False})
        self.assertIn("Page 3/3", P.page_texts()[2])
        P.page_numbers({"format": "- {n} -", "start": 10, "position": "top-right", "pages": "1", "skip_first": False})
        self.assertIn("- 10 -", P.page_texts()[0])
        with self.assertRaises(P.PdfError):
            P.watermark({"text": ""})

    def test_compress(self):
        import io
        import random
        from PIL import Image
        random.seed(1)
        im = Image.effect_noise((1800, 1800), 60).convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "PNG")
        d = pymupdf.open()
        p = d.new_page()
        p.insert_image(pymupdf.Rect(0, 0, 400, 400), stream=buf.getvalue())
        path = os.path.join(self.tmp, "big.pdf")
        d.save(path)
        P.add_files([(path, "big.pdf")])
        self.assertTrue(P.compress("medium")["kept"])
        P.undo()
        res = P.compress("strong")
        self.assertTrue(res["kept"])
        self.assertLess(res["after"], res["before"] * 0.5)
        again = P.compress("light")
        self.assertFalse(again["kept"] and again["after"] > res["after"])


class Security(Base):
    def test_protect_and_unlock(self):
        self.open_a()
        out = P.protect("ouvre", "", allow_print=False, allow_copy=False, allow_modify=True)[0]
        d = pymupdf.open(out)
        self.assertTrue(d.needs_pass)
        self.assertFalse(d.authenticate("mauvais"))
        self.assertTrue(d.authenticate("ouvre"))
        self.assertFalse(d.permissions & pymupdf.PDF_PERM_PRINT)
        self.assertFalse(d.permissions & pymupdf.PDF_PERM_COPY)
        with self.assertRaises(P.PdfError):
            P.protect("", "", True, True, True)
        P.new_session()
        with self.assertRaises(P.PdfError) as cm:
            P.add_files([(out, "p.pdf")])
        self.assertTrue(str(cm.exception).startswith("PASSWORD:"))
        P.add_files([(out, "p.pdf")], password="ouvre")
        self.assertIn("Alpha un", self.texts()[0])
        self.assertFalse(pymupdf.open(P.S["path"]).needs_pass)

    def test_restrictions_only(self):
        self.open_a()
        out = P.protect("", "proprio", allow_print=False, allow_copy=True, allow_modify=False)[0]
        d = pymupdf.open(out)
        self.assertFalse(d.needs_pass)
        self.assertFalse(d.permissions & pymupdf.PDF_PERM_PRINT)


class Convert(Base):
    def test_exports(self):
        self.open_a()
        for fmt in ("docx", "xlsx", "pptx"):
            out = P.convert_from_pdf(fmt)[0]
            self.assertTrue(os.path.getsize(out) > 2000, fmt)
        from pptx import Presentation
        prs = Presentation(out)
        self.assertEqual(len(prs.slides), 3)
        self.assertTrue(any("Alpha un" in sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame))

    @unittest.skipUnless(P.soffice_exe(), "LibreOffice absent")
    def test_office_to_pdf(self):
        src = os.path.join(self.tmp, "note.txt")
        open(src, "w", encoding="utf-8").write("Bonjour LibreOffice\n")
        P.add_files([(src, "note.txt")])
        self.assertIn("Bonjour LibreOffice", self.texts()[0])

    def test_ocr_without_tesseract_is_a_clear_error(self):
        self.open_a()
        if not P.ocr_available():
            with self.assertRaises(P.PdfError) as cm:
                P.ocr()
            self.assertIn("Tesseract", str(cm.exception))
        with self.assertRaises(P.PdfError):
            P.ocr("fra; rm -rf")


class Ai(Base):
    def setUp(self):
        super().setUp()
        import types
        self.calls = []
        fake = types.ModuleType("nasko_llm")

        class LLMError(Exception):
            pass
        fake.LLMError = LLMError
        fake.configured = lambda cfg: True

        def complete(cfg, system, user, max_tokens=160, clean=True):
            self.calls.append(user)
            if "JSON" in system:
                import json
                obj = json.loads(user)
                return json.dumps({k: "[EN] " + v for k, v in obj.items()})
            return "réponse (p. 1)"
        fake.complete = complete
        self._old = sys.modules.get("nasko_llm")
        sys.modules["nasko_llm"] = fake

    def tearDown(self):
        if self._old is not None:
            sys.modules["nasko_llm"] = self._old
        else:
            sys.modules.pop("nasko_llm", None)
        super().tearDown()

    def test_chat_summary_translate(self):
        self.open_a()
        self.assertEqual(P.chat("Que dit la page 2 ?"), "réponse (p. 1)")
        self.assertIn("Alpha deux", self.calls[-1])
        self.assertEqual(P.summarize("short"), "réponse (p. 1)")
        res = P.translate("en")
        d = pymupdf.open(res["files"][0])
        self.assertIn("[EN] Alpha un", d[0].get_text())
        self.assertNotIn("Alpha un\n", d[0].get_text().replace("[EN] Alpha un", ""))
        with self.assertRaises(P.PdfError):
            P.translate("xx")

    def test_chat_picks_relevant_pages_when_long(self):
        long = [("remplissage " * 400) for _ in range(12)]
        long[7] = "Le contrat se termine le 31 décembre 2027. " + long[7]
        ctx, partial = P._context_for("Quelle est la date de fin du contrat ?", long)
        self.assertTrue(partial)
        self.assertIn("Page 8", ctx)
        self.assertLessEqual(len(ctx), P.CHAT_BUDGET + 600)

    def test_scanned_document_asks_for_ocr(self):
        d = pymupdf.open()
        d.new_page()
        path = os.path.join(self.tmp, "blank.pdf")
        d.save(path)
        P.add_files([(path, "blank.pdf")])
        with self.assertRaises(P.PdfError) as cm:
            P.chat("quoi ?")
        self.assertIn("OCR", str(cm.exception))


class Jobs(Base):
    def test_job_lifecycle(self):
        import time
        self.open_a()
        P.start_job("compression", lambda: P.compress("light"))
        for _ in range(100):
            if not P.job()["busy"]:
                break
            time.sleep(0.05)
        j = P.job()
        self.assertFalse(j["busy"])
        self.assertEqual(j["error"], "")
        self.assertIn("before", j["result"])
        P.start_job("boom", lambda: (_ for _ in ()).throw(P.PdfError("raté")))
        for _ in range(100):
            if not P.job()["busy"]:
                break
            time.sleep(0.05)
        self.assertEqual(P.job()["error"], "raté")


if __name__ == "__main__":
    unittest.main()
