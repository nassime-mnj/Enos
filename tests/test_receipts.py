"""Notes de frais (nasko_receipts). Lancer : python tests/test_receipts.py"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import types
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import nasko_receipts as R  # noqa: E402
from PIL import Image  # noqa: E402


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        R.ROOT = self.tmp
        R.UP = os.path.join(self.tmp, ".up")
        R.OUT = os.path.join(self.tmp, "out")
        R.JOB.update(busy=False, result=None, error="")
        R.ensure_dirs()
        self.replies = []
        self.seen = []
        fake = types.ModuleType("nasko_llm")

        class LLMError(Exception):
            pass
        fake.LLMError = LLMError
        fake.configured = lambda cfg: True

        def vision_raw(cfg, system, ask, jpeg, max_tokens=600):
            self.seen.append(jpeg)
            r = self.replies.pop(0)
            if isinstance(r, Exception):
                raise r
            return r if isinstance(r, str) else "Voici : " + json.dumps(r)
        fake.vision_raw = vision_raw
        self._old = sys.modules.get("nasko_llm")
        sys.modules["nasko_llm"] = fake
        self.LLMError = LLMError

    def tearDown(self):
        if self._old is not None:
            sys.modules["nasko_llm"] = self._old
        else:
            sys.modules.pop("nasko_llm", None)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def img(self, name="t.jpg", size=(300, 500)):
        p = os.path.join(self.tmp, name)
        Image.new("RGB", size, (240, 240, 240)).save(p)
        return p

    def deposit(self, *paths):
        ref = R.new_ref()
        for i, p in enumerate(paths):
            shutil.copy(p, os.path.join(R.UP, ref, f"{i:03d}_{os.path.basename(p)}"))
        return ref

    def wait(self):
        for _ in range(200):
            if not R.job()["busy"]:
                break
            time.sleep(0.03)
        return R.job()


class Parsing(Base):
    def test_numbers(self):
        for raw, exp in [(12.5, 12.5), ("12,50 €", 12.5), ("1 234,56", 1234.56), ("1.234,56", 1234.56), ("1,234.56", 1234.56), ("EUR 8", 8.0), (None, None), ("abc", None), (True, None), ("-", None)]:
            self.assertEqual(R._num(raw), exp, raw)

    def test_clean_row_flags_doubtful_data(self):
        ok = R.clean_row({"merchant": " Boulangerie  Paul ", "date": "2025-03-04", "total": "12,60", "vat": 1.15, "currency": "eur", "category": "restauration", "invoice_number": "A-12"}, "a.jpg")
        self.assertEqual((ok["merchant"], ok["date"], ok["total"], ok["category"], ok["currency"], ok["warn"]), ("Boulangerie Paul", "2025-03-04", 12.6, "Restauration", "EUR", []))
        bad = R.clean_row({"merchant": "", "date": "31/02/2025", "total": None, "vat": 5}, "b.jpg")
        self.assertEqual(sorted(bad["warn"]), ["Commerçant absent.", "Date absente ou illisible.", "Montant total absent ou illisible."])
        self.assertIn("TVA incohérente avec le total.", R.clean_row({"merchant": "X", "date": "2025-03-04", "total": 10, "vat": 12}, "c")["warn"])
        self.assertIn("Date improbable.", R.clean_row({"merchant": "X", "date": "1999-01-01", "total": 5}, "d")["warn"])
        self.assertIn("illisible", R.clean_row({"readable": False}, "e")["warn"][0])
        self.assertEqual(R.clean_row({"merchant": "X", "category": "N'importe quoi", "total": 1}, "f")["category"], "Autre")

    def test_to_jpeg_image_and_pdf(self):
        j = R.to_jpeg(self.img(size=(4000, 3000)))
        self.assertEqual(j[:2], b"\xff\xd8")
        self.assertLessEqual(max(Image.open(__import__("io").BytesIO(j)).size), 1800)
        import pymupdf
        d = pymupdf.open()
        d.new_page().insert_text((50, 80), "TICKET 12,60")
        p = os.path.join(self.tmp, "t.pdf")
        d.save(p)
        self.assertEqual(R.to_jpeg(p)[:2], b"\xff\xd8")
        with self.assertRaises(R.ReceiptError):
            R.to_jpeg(os.path.join(self.tmp, "x.png"))


class Flow(Base):
    def test_batch_reads_each_ticket_and_deletes_photos(self):
        a, b, c = self.img("a.jpg"), self.img("b.jpg"), self.img("c.jpg")
        self.replies = [{"merchant": "Total", "date": "2025-05-02", "total": 58.4, "vat": 9.73, "category": "Carburant"},
                        "pas du JSON", "toujours pas",                                              # 2 essais perdus -> ligne à vérifier
                        {"merchant": "SNCF", "date": "2025-05-03", "total": "45,00", "category": "Transport"}]
        ref = self.deposit(a, b, c)
        R.start(ref)
        j = self.wait()
        self.assertEqual(j["error"], "")
        rows = j["result"]["rows"]
        self.assertEqual([(r["file"], r["total"]) for r in rows], [("a.jpg", 58.4), ("b.jpg", None), ("c.jpg", 45.0)])
        self.assertTrue(rows[1]["warn"] and "illisible" in rows[1]["warn"][0])
        self.assertFalse(os.path.exists(os.path.join(R.UP, ref)))                                    # photos supprimées après lecture

    def test_global_failure_stops_early(self):
        ref = self.deposit(self.img("a.jpg"), self.img("b.jpg"))
        self.replies = [self.LLMError("LLM : clé API refusée"), self.LLMError("jamais appelé")]
        R.start(ref)
        j = self.wait()
        self.assertIn("clé API refusée", j["error"])
        self.assertEqual(len(self.replies), 1)                                                       # le 2e ticket n'a pas été tenté


class Export(Base):
    def test_xlsx_formulas_compute(self):
        rows = [{"date": "2025-05-02", "merchant": "Total", "category": "Carburant", "invoice": "F1", "total": 58.4, "vat": 9.73, "currency": "EUR", "file": "a.jpg"},
                {"date": "2025-05-03", "merchant": "SNCF", "category": "Transport", "invoice": "", "total": "45,00", "vat": None, "currency": "EUR", "file": "b.jpg"},
                {"date": "", "merchant": "Resto", "category": "Restauration", "invoice": "", "total": 22.1, "vat": 2.01, "currency": "EUR", "file": "c.jpg"}]
        path = R.export_xlsx(rows)
        import openpyxl
        ws = openpyxl.load_workbook(path)["Notes de frais"]
        self.assertEqual([c.value for c in ws[1]][:3], ["Date", "Commerçant", "Catégorie"])
        self.assertEqual(ws["E3"].value, 45.0)
        self.assertEqual(ws["E6"].value, "=SUM(E2:E4)")
        # recalcul réel par LibreOffice : les formules donnent les bons totaux
        so = shutil.which("soffice")
        if so:
            out = os.path.join(self.tmp, "csv")
            subprocess.run([so, f"-env:UserInstallation=file://{self.tmp}/lo", "--headless", "--convert-to", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1", "--outdir", out, path], capture_output=True, timeout=120)
            files = sorted(os.listdir(out))
            tot = open(os.path.join(out, [f for f in files if "Notes" in f][0]), encoding="utf-8").read()
            syn = open(os.path.join(out, [f for f in files if "Synth" in f][0]), encoding="utf-8").read()
            self.assertIn("125.5", tot.replace(",,", ",").replace('"', ""))
            self.assertIn("Carburant,1,58.4,9.73", syn.replace('"', ""))
            self.assertIn("TOTAL,3,125.5,11.74", syn.replace('"', ""))
        with self.assertRaises(R.ReceiptError):
            R.export_xlsx([])


if __name__ == "__main__":
    unittest.main()
